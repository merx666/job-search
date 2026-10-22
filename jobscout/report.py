from __future__ import annotations
import datetime as dt
import json
from pathlib import Path
from .core import PLATFORMS, digest, safe_url
from .workflow import write_actions
ROOT = Path(__file__).resolve().parent.parent

def spreadsheet_text(value):
    s = str(value or '')
    return "'" + s if s.lstrip().startswith(('=', '+', '-', '@')) else s

def write_report(jobs, plan, statuses, out, demo=False):
    out = Path(out).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    statuses = {k: {**statuses.get(k, {'status':'not_scanned'}),
                    'url': safe_url(statuses.get(k, {}).get('url'), k)} for k in plan['platforms']}
    actions = write_actions(out, plan, jobs, statuses)
    now = dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).isoformat(timespec='seconds')
    meta = {'generated_at':now,'plan':plan,'plan_digest':digest(plan),'platforms':PLATFORMS,
            'statuses':statuses,'demo':demo,'targets':actions,
            'counts':{k:sum(j['bucket']==k for j in jobs) for k in ('shortlist','review','reserve','excluded')}}
    failures = [s for s in statuses.values() if s.get('status') not in ('ok','empty')]
    meta['status'] = 'partial' if failures or not actions['complete'] else 'ok'
    payload = json.dumps({'jobs': jobs, 'meta':meta},ensure_ascii=False).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
    template = (ROOT/'assets/report.html').read_text(encoding='utf-8')
    (out/'jobs.html').write_text(template.replace('__REPORT_DATA__', payload),encoding='utf-8')
    (out/'jobs.json').write_text(json.dumps(jobs,ensure_ascii=False,indent=2),encoding='utf-8')
    (out/'jobs.jsonl').write_text(''.join(json.dumps(j,ensure_ascii=False)+'\n' for j in jobs),encoding='utf-8')
    (out/'summary.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    (out/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError('缺少 openpyxl。HTML/JSON 已保存；运行 pip install -r requirements.txt 后重新生成 Excel。')
    wb=Workbook(); wb.remove(wb.active)
    cols=[('title','岗位'),('company','公司'),('salary_display','薪资原文'),('city','城市'),('experience','经验'),('degree','学历'),('score','排序分（非录用概率）'),('publish_time','发布时间'),('refresh_time','刷新时间'),('hr_active','HR 活跃'),('hr_reply','回复原文'),('platform','平台'),('score_reasons','投递理由'),('unknowns','待核实'),('rejections','排除原因'),('link','岗位链接'),('source_url','来源页'),('captured_at','采集时间')]
    for bucket,title in [('shortlist','优先投递'),('review','待核实'),('reserve','备选'),('excluded','已排除')]:
        ws=wb.create_sheet(title); ws.append([label for _,label in cols])
        for j in jobs:
            if j['bucket'] != bucket: continue
            values=[]
            for key,_ in cols:
                val=j.get(key,'')
                values.append(val if isinstance(val,(int,float)) else spreadsheet_text('；'.join(val) if isinstance(val,list) else val))
            ws.append(values)
            for col,key in [(16,'link'),(17,'source_url')]:
                if j[key]: ws.cell(ws.max_row,col).hyperlink=j[key]
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:
            c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244E43')
        for i in range(1,len(cols)+1): ws.column_dimensions[get_column_letter(i)].width=24 if i<13 else 48
        for row in ws.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    ws=wb.create_sheet('搜索范围与缺口')
    ws.append(['平台','状态','目标数','符合数','待核实数','还差','下一步'])
    for row in actions['platforms']:
        ws.append([row['name'],row['status'],row['target'],row['qualified'],row['review'],row['missing'],spreadsheet_text(row['message'])])
    ws.append(['生成时间',now]); ws.append(['示例数据',str(demo)])
    ws.append(['评分说明','规则排序分，不代表能力评价或录用概率；未知信息不能当作满足条件。'])
    wb.save(out/'jobs.xlsx')
    return meta
